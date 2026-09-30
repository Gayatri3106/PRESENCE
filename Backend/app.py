import base64
import binascii
import hashlib
import hmac
import io
import json
import logging
import os
import secrets
import time
from datetime import datetime, timedelta, timezone

from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from werkzeug.security import check_password_hash
from apscheduler.schedulers.background import BackgroundScheduler
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
import qrcode

from config import config
from face_pipeline import decode_base64_image, get_face_embedding, compare_embeddings
from room_code import generate_room_code
from supabase_client import (
    get_student_by_email,
    get_student_by_id,
    get_student_by_roll_no,
    create_student,
    update_student_embedding,
    get_faculty_by_id,
    get_faculty_by_email,
    get_classes_for_faculty,
    create_session,
    get_active_session,
    get_active_session_for_class,
    close_session,
    get_expired_active_sessions,
    validate_room_code,
    create_attendance_record,
    student_already_checked_in,
    get_records_for_session,
    get_attendance_for_student,
    get_students_by_class,
    get_class_attendance_stats,
    get_session_by_id,
)

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024
cors_origins = config.CORS_ORIGINS.split(",") if config.CORS_ORIGINS and config.CORS_ORIGINS != "*" else "*"
CORS(app, resources={r"/api/*": {"origins": cors_origins}})

ROOM_TTL = 120
QR_TTL = 30


def success(message, data=None, status=200):
    return jsonify({"success": True, "message": message, "data": data}), status


def error(message, status=400, details=None):
    body = {"success": False, "message": message}
    if details is not None:
        body["details"] = details
    return jsonify(body), status


def get_json():
    data = request.get_json(silent=True)
    return data if isinstance(data, dict) else None


def clean_string(value):
    return "" if value is None else str(value).strip()


def parse_iso(value):
    if isinstance(value, datetime):
        dt = value
    else:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def now_utc():
    return datetime.now(timezone.utc)


def qr_secret():
    return config.FLASK_SECRET_KEY


def current_qr_bucket(ts=None):
    ts = ts or now_utc()
    return int(ts.timestamp()) // QR_TTL


def make_qr_token(session_id, bucket=None):
    bucket = current_qr_bucket() if bucket is None else int(bucket)
    payload = f"{session_id}:{bucket}"
    sig = hmac.new(qr_secret().encode(), payload.encode(), hashlib.sha256).hexdigest()
    raw = f"{payload}:{sig}"
    return base64.urlsafe_b64encode(raw.encode()).decode().rstrip("=")


def validate_qr_token(session_id, token):
    try:
        padded = token + "=" * (-len(token) % 4)
        raw = base64.urlsafe_b64decode(padded.encode()).decode()
        token_session, bucket_text, signature = raw.split(":", 2)
        if token_session != str(session_id):
            return False
        bucket = int(bucket_text)
        if bucket != current_qr_bucket():
            return False
        expected = hmac.new(
            qr_secret().encode(),
            f"{token_session}:{bucket}".encode(),
            hashlib.sha256,
        ).hexdigest()
        return hmac.compare_digest(signature, expected)
    except Exception:
        return False


def session_qr_payload(session):
    token = make_qr_token(session["id"])
    expires = datetime.fromtimestamp((current_qr_bucket() + 1) * QR_TTL, tz=timezone.utc)
    qr_value = json.dumps({
        "type": "presence-attendance",
        "session_id": str(session["id"]),
        "room_code": session["room_code"],
        "token": token,
    }, separators=(",", ":"))
    qr_img = qrcode.make(qr_value)
    qr_bytes = io.BytesIO()
    qr_img.save(qr_bytes, format="PNG")
    qr_image = "data:image/png;base64," + base64.b64encode(qr_bytes.getvalue()).decode("ascii")
    return {
        "session_id": session["id"],
        "room_code": session["room_code"],
        "token": token,
        "qr_value": qr_value,
        "qr_image": qr_image,
        "valid_until": expires.isoformat(),
        "valid_for_seconds": max(0, int((expires - now_utc()).total_seconds())),
    }


def session_expired(session):
    return parse_iso(session["expires_at"]) <= now_utc()


def expire_sessions_job():
    try:
        for session in get_expired_active_sessions():
            try:
                close_session(session["id"])
                try:
                    save_attendance_report(session["id"])
                except Exception:
                    logger.exception("Could not automatically generate report for %s", session.get("id"))
                logger.info("Automatically closed expired attendance session %s", session["id"])
            except Exception:
                logger.exception("Could not close session %s", session.get("id"))
    except Exception:
        logger.exception("Session expiry job failed")


# ---------------- health ----------------
@app.get("/api/health")
def health():
    return success("OK", {"status": "alive", "time": now_utc().isoformat()})


# ---------------- student registration ----------------
@app.post("/api/register-face")
def register_face():
    start = time.perf_counter()
    try:
        data = get_json()
        if not data:
            return error("Request body is required.")

        name = clean_string(data.get("name"))
        email = clean_string(data.get("email")).lower()
        roll_no = clean_string(data.get("roll_no")).upper()
        image = data.get("image")

        if not name or not email or not roll_no or not image:
            return error("Name, Email ID, Roll Number and live face capture are required.")
        if "@" not in email or "." not in email.rsplit("@", 1)[-1]:
            return error("Please enter a valid email address.")

        existing_roll = get_student_by_roll_no(roll_no)
        existing_email = get_student_by_email(email)

        if existing_roll and existing_roll.get("face_embedding"):
            return error("This Roll Number is already registered.", 409)
        if existing_email and existing_email.get("face_embedding"):
            return error("This Email ID is already registered.", 409)
        if existing_roll and existing_email and existing_roll.get("id") != existing_email.get("id"):
            return error("Roll Number and Email ID belong to different students.", 409)

        student = existing_roll or existing_email
        if not student:
            student = create_student(name=name, email=email, roll_no=roll_no)
        else:
            if student.get("name") != name or student.get("email") != email:
                from supabase_client import update_student_identity
                update_student_identity(student["id"], name, email, roll_no)

        try:
            image_array = decode_base64_image(image)
            embedding = get_face_embedding(image_array)
        except (ValueError, binascii.Error) as exc:
            return error(str(exc))

        if len(embedding) != 128:
            return error("Face embedding generation failed: expected 128 dimensions.", 500)

        update_student_embedding(student["id"], [float(x) for x in embedding])
        return success("Face registered successfully.", {
            "student_id": student["id"],
            "roll_no": roll_no,
            "embedding_dim": 128,
            "compute_time_s": round(time.perf_counter() - start, 3),
        }, 201)
    except Exception as exc:
        logger.exception("Face registration failed")
        return error("Face registration failed.", 500, str(exc))


# ---------------- teacher/student authentication ----------------
@app.post("/api/student/login")
def student_login():
    try:
        data = get_json() or {}
        roll_no = clean_string(data.get("roll_no")).upper()
        email = clean_string(data.get("email")).lower()
        if not roll_no or not email:
            return error("Roll Number and Email ID are required.")
        student = get_student_by_roll_no(roll_no)
        if not student or str(student.get("email", "")).lower() != email:
            return error("Student details were not found. Complete face registration first.", 401)
        if not student.get("face_embedding"):
            return error("Face is not registered for this student. Please complete registration first.", 400)
        return success("Student login successful.", {
            "student_id": student["id"], "roll_no": student.get("roll_no"),
            "name": student.get("name"), "email": student.get("email")
        })
    except Exception as exc:
        logger.exception("Student login failed")
        return error("Student login failed.", 500, str(exc))


@app.get("/api/room/<room_code>")
def room_lookup(room_code):
    session = validate_room_code(clean_string(room_code).upper())
    if not session:
        return error("Room code is invalid or expired.", 404)
    return success("Active room found.", {
        "id": session["id"], "session_id": session["id"], "room_code": session["room_code"],
        "subject": session.get("subject", ""), "period_start": session.get("period_start", ""),
        "period_end": session.get("period_end", ""), "expires_at": session.get("expires_at"), "status": session.get("status")
    })


@app.post("/api/teacher/login")
def teacher_login():
    try:
        data = get_json() or {}
        email = clean_string(data.get("email")).lower()
        password = str(data.get("password", ""))
        if not email or not password:
            return error("Email and password are required.")

        faculty = get_faculty_by_email(email)
        if not faculty:
            return error("Teacher account not found.", 401)

        stored_hash = faculty.get("password_hash")
        if not stored_hash or not check_password_hash(stored_hash, password):
            return error("Invalid teacher credentials.", 401)

        return success("Teacher login successful.", {
            "faculty_id": faculty["id"],
            "name": faculty.get("name"),
            "email": faculty.get("email"),
            "subject": faculty.get("subject") or faculty.get("department") or "",
            "department": faculty.get("department"),
        })
    except Exception as exc:
        logger.exception("Teacher login failed")
        return error("Teacher login failed.", 500, str(exc))


@app.get("/api/teacher/<faculty_id>")
def teacher_details(faculty_id):
    faculty = get_faculty_by_id(faculty_id)
    if not faculty:
        return error("Teacher not found.", 404)
    return success("Teacher details retrieved.", {
        "faculty_id": faculty["id"],
        "name": faculty.get("name"),
        "email": faculty.get("email"),
        "subject": faculty.get("subject") or faculty.get("department") or "",
        "department": faculty.get("department"),
    })


@app.get("/api/teacher/<faculty_id>/classes")
def teacher_classes(faculty_id):
    faculty = get_faculty_by_id(faculty_id)
    if not faculty:
        return error("Teacher not found.", 404)
    return success("Classes retrieved.", get_classes_for_faculty(faculty_id))


# ---------------- room/session ----------------
@app.post("/api/create-session")
def create_attendance_session():
    try:
        data = get_json() or {}
        faculty_id = clean_string(data.get("faculty_id"))
        class_id = clean_string(data.get("class_id"))
        subject = clean_string(data.get("subject"))
        period_start = clean_string(data.get("period_start"))
        period_end = clean_string(data.get("period_end"))

        if not all([faculty_id, class_id, subject, period_start, period_end]):
            return error("Faculty, class, subject, period start and period end are required.")

        faculty = get_faculty_by_id(faculty_id)
        if not faculty:
            return error("Faculty not found.", 404)

        classes = get_classes_for_faculty(faculty_id)
        selected = next((c for c in classes if str(c.get("id")) == class_id), None)
        if not selected:
            return error("This class is not assigned to the faculty.", 403)

        existing = get_active_session_for_class(class_id)
        if existing and not session_expired(existing):
            return error("An attendance room is already active for this class.", 409, {
                "session_id": existing["id"], "room_code": existing["room_code"], "expires_at": existing["expires_at"]
            })
        if existing:
            close_session(existing["id"])

        room_code = generate_room_code()
        generated_at = now_utc()
        expires_at = generated_at + timedelta(seconds=ROOM_TTL)
        session = create_session(
            class_id=class_id,
            faculty_id=faculty_id,
            subject=subject,
            period_start=period_start,
            period_end=period_end,
            room_code=room_code,
            expires_at=expires_at,
        )

        return success("Attendance room created successfully.", {
            "session_id": session["id"],
            "faculty_id": faculty_id,
            "class_id": class_id,
            "class_name": selected.get("name") or selected.get("class_name") or "",
            "subject": subject,
            "period_start": period_start,
            "period_end": period_end,
            "room_code": room_code,
            "generated_at": generated_at.isoformat(),
            "expires_at": expires_at.isoformat(),
            "valid_for_seconds": ROOM_TTL,
            "status": "active",
        }, 201)
    except Exception as exc:
        logger.exception("Session creation failed")
        return error("Unable to create attendance room.", 500, str(exc))


@app.get("/api/session/<session_id>/qr")
def session_qr(session_id):
    session = get_active_session(session_id)
    if not session:
        return error("Session not found or no longer active.", 404)
    if session_expired(session):
        close_session(session_id)
        return error("Attendance room has expired.", 410)
    return success("Current QR token retrieved.", session_qr_payload(session))


# ---------------- attendance ----------------
@app.post("/api/checkin")
def checkin():
    try:
        data = get_json() or {}
        roll_no = clean_string(data.get("roll_no")).upper()
        room_code = clean_string(data.get("room_code")).upper()
        qr_token = clean_string(data.get("qr_token"))
        session_id = clean_string(data.get("session_id"))
        image = data.get("image")

        missing = []
        if not roll_no: missing.append("roll_no")
        if not room_code: missing.append("room_code")
        if not qr_token: missing.append("qr_token")
        if not session_id: missing.append("session_id")
        if not image: missing.append("image")

        if missing:
            return error(f"Missing required fields: {', '.join(missing)}", 400)

        session = validate_room_code(room_code)
        if not session:
            return error("Room code is invalid or expired.", 401)
        if str(session["id"]) != str(session_id):
            return error("The scanned QR does not belong to this room.", 401)
        if not validate_qr_token(session_id, qr_token):
            return error("QR code is invalid or has expired. Scan the current QR again.", 401)

        student = get_student_by_roll_no(roll_no)
        if not student:
            return error("Student not found. Please complete registration first.", 404)
        stored = student.get("face_embedding")
        if not stored:
            return error("Face is not registered for this Roll Number.", 400)
        if student_already_checked_in(session["id"], student["id"]):
            return error("Attendance is already marked for this room.", 409)

        try:
            live_image = decode_base64_image(image)
            live_embedding = get_face_embedding(live_image)
        except (ValueError, binascii.Error) as exc:
            return error(f"Invalid face image format: {str(exc)}", 400)

        result = compare_embeddings(stored, live_embedding)
        if not result["match"]:
            return error("Face verification failed. Attendance was not marked.", 401, {
                "face_match": False,
                "distance": result["distance"],
                "threshold": result["threshold"],
            })

        record = create_attendance_record(
            session_id=session["id"],
            student_id=student["id"],
            match_distance=result["distance"],
            status="present",
            method="qr+face"
        )
        return success("Attendance marked successfully!", {
            "record_id": record["id"],
            "session_id": session["id"],
            "roll_no": student["roll_no"],
            "status": "present",
            "distance": result["distance"],
            "threshold": result["threshold"],
        })
    except Exception as exc:
        logger.exception("Check-in failed")
        return error("Attendance check-in failed.", 500, str(exc))


# ---------------- session/report ----------------
def get_session_any(session_id):
    return get_session_by_id(session_id)


def build_attendance_workbook(session_id):
    session = get_active_session(session_id) or get_session_any(session_id)
    if not session:
        raise ValueError("Session not found.")

    records = get_records_for_session(session_id)
    students = get_students_by_class(session["class_id"])
    present_by_id = {str(r["student_id"]): r for r in records if r.get("status") == "present"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    faculty = get_faculty_by_id(session.get("faculty_id")) if session.get("faculty_id") else None
    class_name = session.get("class_name") or session.get("class_id") or ""

    meta = [
        ("Teacher Name", (faculty or {}).get("name", "")),
        ("Teacher Email", (faculty or {}).get("email", "")),
        ("Subject", session.get("subject", "")),
        ("Class", class_name),
        ("Period Start", session.get("period_start", "")),
        ("Period End", session.get("period_end", "")),
        ("Room Code", session.get("room_code", "")),
        ("Date", parse_iso(session["generated_at"]).date().isoformat()),
        ("Total Registered", len(students)),
        ("Total Present", len(present_by_id)),
        ("Total Absent", max(0, len(students) - len(present_by_id))),
    ]
    for row, (label, value) in enumerate(meta, start=1):
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)

    header_row = len(meta) + 3
    headers = ["S.No", "Roll Number", "Student Name", "Email", "Status", "Attendance Time"]
    for col, value in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col, value=value)
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    for i, student in enumerate(sorted(students, key=lambda x: str(x.get("roll_no", ""))), 1):
        record = present_by_id.get(str(student["id"]))
        values = [
            i,
            student.get("roll_no", ""),
            student.get("name", ""),
            student.get("email", ""),
            "PRESENT" if record else "ABSENT",
            record.get("timestamp", "") if record else "",
        ]
        for col, value in enumerate(values, 1):
            ws.cell(row=header_row + i, column=col, value=value)

    widths = [8, 18, 26, 34, 14, 28]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def save_attendance_report(session_id):
    reports_dir = os.path.join(os.path.dirname(__file__), "reports")
    os.makedirs(reports_dir, exist_ok=True)
    session = get_session_any(session_id)
    if not session:
        raise ValueError("Session not found.")
    workbook = build_attendance_workbook(session_id)
    date_text = parse_iso(session["generated_at"]).date().isoformat()
    path = os.path.join(reports_dir, f"Attendance_{session.get('room_code', session_id)}_{date_text}.xlsx")
    with open(path, "wb") as f:
        f.write(workbook.getvalue())
    return path


@app.post("/api/close-session")
def close_attendance_session():
    data = get_json() or {}
    session_id = clean_string(data.get("session_id"))
    if not session_id:
        return error("Session ID is required.")
    session = get_session_any(session_id)
    if not session:
        return error("Session not found.", 404)
    close_session(session_id)
    return success("Attendance room closed and report finalized.", {"session_id": session_id, "status": "closed"})


@app.get("/api/session/<session_id>")
def get_session(session_id):
    session = get_session_any(session_id)
    if not session:
        return error("Session not found.", 404)
    if session.get("status") == "active" and session_expired(session):
        close_session(session_id)
        session["status"] = "closed"
    return success("Session retrieved.", session)


@app.get("/api/session/<session_id>/records")
def session_records(session_id):
    return success("Attendance records retrieved.", {"session_id": session_id, "records": get_records_for_session(session_id)})


@app.get("/api/session/<session_id>/export")
def export_session(session_id):
    try:
        session = get_session_any(session_id)
        if not session:
            return error("Session not found.", 404)
        if session.get("status") == "active" and session_expired(session):
            close_session(session_id)
        workbook = build_attendance_workbook(session_id)
        date_text = parse_iso(session["generated_at"]).date().isoformat()
        filename = f"Attendance_{session.get('room_code', session_id)}_{date_text}.xlsx"
        return send_file(workbook, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=filename)
    except Exception as exc:
        logger.exception("Excel export failed")
        return error("Unable to generate attendance Excel file.", 500, str(exc))


@app.get("/api/student/<student_id>/attendance")
def student_attendance(student_id):
    student = get_student_by_id(student_id)
    if not student:
        return error("Student not found.", 404)
    return success("Attendance retrieved.", {"student": {"id": student["id"], "name": student.get("name"), "email": student.get("email"), "roll_no": student.get("roll_no")}, "records": get_attendance_for_student(student_id)})


@app.get("/api/student/<student_id>/percentage")
def student_percentage(student_id):
    records = get_attendance_for_student(student_id)
    total = len(records)
    present = sum(1 for r in records if r.get("status") == "present")
    return success("Attendance percentage calculated.", {"student_id": student_id, "present": present, "total": total, "percentage": round(present * 100 / total, 1) if total else 0})


@app.get("/api/class/<class_id>/stats")
def class_stats(class_id):
    return success("Class attendance statistics retrieved.", {"class_id": class_id, "students": get_class_attendance_stats(class_id)})


scheduler = BackgroundScheduler(daemon=True)
scheduler.add_job(expire_sessions_job, "interval", seconds=10, id="expire-sessions", replace_existing=True)
scheduler.start()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=5000, debug=config.FLASK_DEBUG, use_reloader=False)